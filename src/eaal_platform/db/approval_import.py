"""Reading a CSV or Excel file of people into rows for ``approvals.import_entries``.

The expected layout is one header row, then one person per row:

    teachers:  name, email
    students:  name, email, enrollment_no

Column names are matched loosely ("E-mail", "Full name", "Roll no" ...), the
order doesn't matter, and extra columns are ignored. A file without a header
row also works if each row contains an email address.
"""

from __future__ import annotations

import csv
import io
from typing import Any

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000

_NAME = {"name", "full name", "fullname", "student name", "teacher name", "professor name",
         "display name"}  # fmt: skip
_EMAIL = {"email", "e-mail", "email address", "email id", "mail", "university email",
          "student email", "teacher email", "professor email"}  # fmt: skip
_ENROLLMENT = {"enrollment_no", "enrollment no", "enrollment", "enrollment number",
               "enrolment_no", "enrolment no", "enrolment", "enrolment number",
               "roll no", "roll number", "roll", "student id", "id"}  # fmt: skip

TEMPLATES = {
    "student": (
        "name,email,enrollment_no\n"
        "Asha Patel,asha.patel@university.edu,BSC-101\n"
        "Ravi Kumar,ravi.kumar@university.edu,BSC-102\n"
    ),
    "professor": (
        "name,email\n"
        "Dr. Meera Shah,meera.shah@university.edu\n"
        "Dr. Arjun Rao,arjun.rao@university.edu\n"
    ),
}


class UploadError(ValueError):
    """The file can't be read; the message tells the administrator what to do."""


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Excel stores 101 as 101.0
    return str(value).strip()


def _csv_rows(data: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeError:
            continue
    else:  # pragma: no cover - cp1252 decodes almost anything
        raise UploadError("That file's text encoding isn't supported. Save it as CSV (UTF-8).")
    first_line = text.splitlines()[0] if text.strip() else ""
    delimiter = max((",", ";", "\t"), key=first_line.count)
    return [
        [_text(cell) for cell in row] for row in csv.reader(io.StringIO(text), delimiter=delimiter)
    ]


def _xlsx_rows(data: bytes) -> list[list[str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:  # pragma: no cover - installed with the server extra
        raise UploadError("Excel files aren't supported here. Save the file as CSV.") from None
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        return [[_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    except Exception as exc:
        raise UploadError("That Excel file couldn't be read. Is it a valid .xlsx file?") from exc


def _grid_from_file(filename: str, data: bytes) -> list[list[str]]:
    """The file as rows of text cells (CSV or Excel), with empty rows dropped."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise UploadError("That file is too large (the limit is 5 MB).")
    if not data:
        raise UploadError("That file is empty.")
    lower = filename.lower()
    if lower.endswith((".xlsx", ".xlsm")):
        grid = _xlsx_rows(data)
    elif lower.endswith(".xls"):
        raise UploadError("Old .xls files aren't supported. Save it as .xlsx or .csv.")
    elif lower.endswith((".csv", ".txt", ".tsv")):
        grid = _csv_rows(data)
    else:
        raise UploadError("Upload a .csv or .xlsx file.")
    grid = [row for row in grid if any(cell for cell in row)]
    if not grid:
        raise UploadError("That file has no rows.")
    return grid


def _cell(cells: list[str], column: int | None) -> str:
    return cells[column] if column is not None and column < len(cells) else ""


def _find(header: list[str], names: Any) -> int | None:
    return next((i for i, h in enumerate(header) if h in names), None)


def _layout(grid: list[list[str]], role: str) -> tuple[list[list[str]], int, dict[str, int | None]]:
    """Work out where the columns are: ``(data rows, number of the first one, columns)``."""
    header = [cell.lower().strip() for cell in grid[0]]
    if any(cell in _EMAIL for cell in header):
        columns = {
            "name": _find(header, _NAME),
            "email": _find(header, _EMAIL),
            "enrollment": _find(header, _ENROLLMENT),
        }
        return grid[1:], 2, columns
    if any("@" in cell for cell in grid[0]):
        # No header row: find the email in each row; the other cells are name, then enrolment.
        return grid, 1, {"name": None, "email": None, "enrollment": None}
    raise UploadError(
        "Couldn't find an Email column. The first row must name the columns "
        "(name, email" + (", enrollment_no" if role == "student" else "") + "). "
        "Download the template to see the layout."
    )


def _person(cells: list[str], columns: dict[str, int | None]) -> tuple[str, str, str]:
    """``(name, email, enrollment)`` from one row."""
    if columns["email"] is not None:
        return (
            _cell(cells, columns["name"]),
            _cell(cells, columns["email"]),
            _cell(cells, columns["enrollment"]),
        )
    at = next((i for i, c in enumerate(cells) if "@" in c), None)
    others = [c for i, c in enumerate(cells) if i != at and c]
    return (
        others[0] if others else "",
        cells[at] if at is not None else "",
        others[1] if len(others) > 1 else "",
    )


def read_rows(filename: str, data: bytes, role: str) -> list[dict[str, Any]]:
    """Parse an uploaded file into ``[{row, name, email, enrollment_no}, ...]``."""
    body, first_number, columns = _layout(_grid_from_file(filename, data), role)
    if len(body) > MAX_ROWS:
        raise UploadError(f"That file has more than {MAX_ROWS} rows. Split it into smaller files.")
    rows: list[dict[str, Any]] = []
    for offset, cells in enumerate(body):
        name, email, enrollment = _person(cells, columns)
        rows.append(
            {
                "row": first_number + offset,
                "name": name,
                "email": email,
                "enrollment_no": enrollment if role == "student" else "",
            }
        )
    return rows
